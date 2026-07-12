from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).with_name('build_report.ps1').read_text(encoding='utf-8')
CONTENT = re.search(r"\$content = @'\n(.*?)\n'@", HERE, re.S).group(1)
TEMPLATE = Path(__file__).with_name('附录A-专业综合实验报告-模版.docx')
OUT = Path(__file__).with_name('飞跃雷区四旋翼飞控系统设计报告_终版.docx')
FIGURES = {
    '[图：fig_height_estimation.svg]': Path(__file__).with_name('figures_png') / 'fig_height_estimation.png',
    '[图：fig_vertical_state.svg]': Path(__file__).with_name('figures_png') / 'fig_vertical_state.png',
    '[图：fig_attitude_debug.svg]': Path(__file__).with_name('figures_png') / 'fig_attitude_debug.png',
}

doc = Document(TEMPLATE)
body = doc._element.body
for child in list(body):
    if not child.tag.endswith('sectPr'):
        body.remove(child)

section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(3.0)
section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')
style.font.size = Pt(12)

# Template-compatible page footer with a real PAGE field.
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
run._r.extend([fld_begin, instr, fld_end])

title_lines = {'中国矿业大学', '专业综合实验', '面向全国大学生智能车大赛飞跃雷区任务的', '四旋翼飞行控制系统设计与实现'}
for raw in CONTENT.splitlines():
    text = raw.strip()
    if not text:
        continue
    if text == '[分页]':
        doc.add_page_break()
        continue
    p = doc.add_paragraph()
    if text in FIGURES and FIGURES[text].exists():
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FIGURES[text]), width=Cm(15.2))
        continue
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')
    run.font.size = Pt(12)
    if text in title_lines:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.bold = True
        run.font.size = Pt(20)
        p.paragraph_format.space_after = Pt(20)
    elif re.match(r'^(摘  要|目  录|参考文献|附录[0-9]|小组周志)$', text):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.bold = True
        run.font.size = Pt(16)
        p.paragraph_format.space_before = Pt(14)
    elif re.match(r'^[1-6] (实验目的与意义|基础知识与理论|系统总体设计|飞行控制系统实现|实验测试与结果分析|总结与展望)', text):
        run.bold = True
        run.font.size = Pt(16)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(9)
    elif re.match(r'^[1-6]\.[0-9] ', text):
        run.bold = True
        run.font.size = Pt(14)
        p.paragraph_format.space_before = Pt(12)
    elif text.startswith('图') or text.startswith('[图：'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.font.size = Pt(10.5)
    elif not text.startswith(('关键词：', '[', '（')):
        p.paragraph_format.first_line_indent = Cm(0.74)

# Engineering evidence appendix: keep extracts short and explain their role.
doc.add_page_break()
for heading, body_text in [
    ('4.6 姿态双环的控制逻辑与调试顺序', '姿态控制不是直接把遥控量送到电机，而是先由角度外环计算期望角速度，再由角速度内环根据陀螺仪反馈产生Roll、Pitch、Yaw控制量。内环负责抑制快速扰动，外环决定机体最终回到哪个姿态。调试时先确认电机编号、旋向和正负方向，再在小油门条件下调节角速度环，最后加入角度环。若跳过内环直接提高角度环增益，飞行器容易出现高频抖动，且难以判断问题来自传感器方向还是PID参数。'),
    ('4.7 高度环的状态估计、轨迹与闭环关系', '高度控制由“目标高度—轨迹生成—目标垂直速度—油门输出”构成。TOF原始距离不直接参与油门控制：它先与IMU世界Z轴加速度共同进入高度估计器，得到平滑的高度z和垂直速度vz。随后高度外环将位置误差转换成受限的速度目标，速度内环再调整油门。这样做能够避免TOF单次跳变直接造成油门突变，也使起飞和到达目标高度时具有可控的加减速过程。'),
    ('4.8 光流估计与水平双环的耦合关系', '光流模块首先给出与高度和地面纹理相关的相对位移。程序使用TOF高度完成比例换算，完成X/Y映射后再补偿探头偏心和姿态旋转引起的假位移。位置外环在地球系计算位置误差并生成速度目标；速度误差必须根据Yaw角旋转到机体系，才能正确映射为Pitch和Roll指令。该转换是后续调试的关键：机体系用于标定光流本体，地球系用于导航和控制，两者不能混用。'),
    ('4.9 位置保持切入与保护逻辑', '位置保持并非解锁后立即开启。程序要求飞行器已解锁、光流有效、高度达到使能阈值、高度环进入保持阶段且垂直/水平速度足够小。满足条件后，系统先捕获当前点作为保持目标，清除速度PID的积分残留，并通过制动阶段和渐入权重逐步增加控制输出。这样可以避免起飞阶段的漂移被积分项放大，也避免刚切换位置环时出现突然大倾角。'),
    ('5.6 基于日志的调试闭环', '本项目的参数修改均以日志为依据，而不是凭飞行体感判断。高度环记录TOF、估计高度、垂直速度、创新量和油门；光流部分记录原始位移、补偿量、机体系/地球系状态；控制部分记录目标值、反馈值和电机输出。每次修改遵循“复现问题—定位状态量—只改一个环节—独立复测”的顺序。该过程使姿态、定高和位置控制能够逐层建立，而不是在多个未稳定环路上同时调参。'),
]:
    p = doc.add_paragraph(); r = p.add_run(heading); r.bold = True; r.font.size = Pt(14)
    p = doc.add_paragraph(body_text); p.paragraph_format.first_line_indent = Cm(0.74); p.paragraph_format.line_spacing = 1.5

summary = Path(__file__).with_name('figures_png') / 'fig_test_summary.png'
if summary.exists():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(summary), width=Cm(15.2))
    p = doc.add_paragraph('图5-3 高度控制与原路返回测试结果汇总'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(10.5)
    p = doc.add_paragraph('图中比较了TOF定高误差、机体系原路返回残差和地球系原路返回残差。机体系残差较小，说明光流本体的方向映射和短时比例标定基本正确；地球系Y方向残差相对较大，反映出航向角误差在坐标旋转中的累计影响。该结果也说明后续应引入视觉信标等绝对观测，修正长期定位漂移。')
    p.paragraph_format.first_line_indent = Cm(0.74); p.paragraph_format.line_spacing = 1.5

# Additional visual evidence and control-logic diagrams.
visuals = [
    ('图4-3 姿态角—角速度级联控制与电机混控流程', 'fig_attitude_control.png', '姿态外环根据Roll、Pitch、Yaw角度误差给出期望角速度，内环利用陀螺仪反馈形成力矩控制量，最终与总油门共同进入四电机混控。'),
    ('图4-4 高度EKF及高度—垂直速度级联控制流程', 'fig_altitude_control.png', 'TOF与IMU世界Z轴加速度共同形成高度和垂直速度估计；轨迹生成器限制目标变化，高度外环给出目标垂直速度，速度内环调节油门。'),
    ('图4-5 光流处理及位置—速度级联控制流程', 'fig_position_control.png', 'LC302原始光流经方向映射、偏心与姿态补偿后得到机体系速度；位置误差在地球系计算，速度误差旋转到机体系后生成Roll与Pitch目标。'),
    ('图4-6 位置保持模式切入与释放状态机', 'fig_state_machine.png', '状态机通过解锁、光流有效性、高度、垂直速度和水平速度条件控制位置环切入；制动、目标捕获和渐入逻辑用于抑制模式切换冲击。'),
    ('图5-4 原路返回闭环测试轨迹示意', 'fig_return_trajectory.png', '原路返回测试优先评价机体系闭环残差，可减少不同地面纹理引起的比例差异；地球系结果同时反映航向误差的累计影响。'),
]
for caption, filename, explanation in visuals:
    img = Path(__file__).with_name('figures_png') / filename
    if img.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Cm(15.2))
        p = doc.add_paragraph(caption); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs: p.runs[0].font.size = Pt(10.5)
        p = doc.add_paragraph(explanation)
        p.paragraph_format.first_line_indent = Cm(0.74); p.paragraph_format.line_spacing = 1.5

doc.add_page_break()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('附录3 关键控制代码摘录'); r.bold = True; r.font.size = Pt(16)
for title, note, image in [
    ('A. 水平位置—速度级联控制', '该段代码先根据位置误差生成期望速度，再将地球系速度误差转换至机体系，输出俯仰和横滚指令，并配合限幅、渐入和制动逻辑，避免位置保持模式切换时产生突变。', 'code_location_cascade.png'),
    ('B. 油门斜坡与电机混控', '该段代码负责将总油门与Roll、Pitch、Yaw控制量组合为四路电机输出。通过油门斜坡和输出边界处理，降低解锁、起飞和姿态修正时的电机指令突变。', 'code_motor_mixing.png'),
]:
    p = doc.add_paragraph(); rr = p.add_run(title); rr.bold = True; rr.font.size = Pt(14)
    p = doc.add_paragraph(note); p.paragraph_format.first_line_indent = Cm(0.74); p.paragraph_format.line_spacing = 1.5
    img = Path(__file__).with_name('code_png') / image
    if img.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(img), width=Cm(15.2))

def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for rr in para.runs:
                    rr.font.name = 'Times New Roman'
                    rr._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                    rr.font.size = Pt(12)

def insert_table_after(anchor_text, rows):
    anchor = next((p for p in doc.paragraphs if p.text.strip() == anchor_text), None)
    if anchor is None:
        return
    table = doc.add_table(rows=len(rows), cols=2)
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
    format_table(table)
    anchor._p.addnext(table._tbl)

insert_table_after('完成日期：2026年7月', [
    ('学院', '信息与控制工程学院'), ('专业', '电子信息工程'),
    ('成员', '（待填写）'), ('班级及学号', '（待填写）'), ('指导教师', '（待填写）')])
insert_table_after('专业综合实验任务书', [
    ('课题名称', '面向全国大学生智能车大赛飞跃雷区任务的四旋翼飞行控制系统设计与实现'),
    ('主要内容及基本要求', '完成姿态解算、TOF高度估计、LC302光流处理、姿态/高度/位置级联控制、电机混控、日志分析及视觉信标接口预留。')])
insert_table_after('成员个人工作总结', [
    ('姓名、班级、学号', '（待填写）'),
    ('本人分工', '飞控软件、传感器融合、控制算法、参数调试与日志分析'),
    ('涉及知识与工具', '嵌入式C、Mahony姿态解算、EKF、级联PID、TOF与光流定位'),
    ('完成情况与自评', '完成姿态、定高和短时位置保持基础功能；个人自评（待填写）')])

doc.save(OUT)
print(OUT)
